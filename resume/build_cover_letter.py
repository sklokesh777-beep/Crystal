#!/usr/bin/env python3
"""
Builds the cover letter in PDF, DOCX and TXT, reusing the resume's typography so
the two documents read as a matched set.

  S-K-Lokesh-Cover-Letter.pdf   - attach this to the application
  S-K-Lokesh-Cover-Letter.docx  - editable Word master
  S-K-Lokesh-Cover-Letter.txt   - plain text for pasting into a web form or email

Usage:  python3 build_cover_letter.py
"""

import os

from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Spacer)

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_resume as R
from build_resume import (CONTACT_SIZE, FONT, NAME_SIZE, P, Rule, _add_hyperlink,
                          _add_runs, _bottom_border, _set_spacing, plain,
                          register_fonts, smart)
from resume_content import CONTACT, NAME
from cover_letter_content import (CLOSING, DATE, ENCLOSURE, PARAGRAPHS,
                                 RECIPIENT, SALUTATION, SIGNOFF, SUBJECT)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_BASE = os.path.join(HERE, "S-K-Lokesh-Cover-Letter")

BODY = 10.2          # a letter can breathe more than a resume
LEAD = 13.4
MARGIN_X = 20.0 * mm
MARGIN_TOP = 13.0 * mm
MARGIN_BOTTOM = 13.0 * mm
FRAME_W = A4[0] - 2 * MARGIN_X


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def build_pdf(path=OUT_BASE + ".pdf"):
    register_fonts()
    body = dict(fontName=FONT, fontSize=BODY, leading=LEAD, textColor="#000000")
    st = {
        "name": ParagraphStyle("n", fontName=FONT + "-Bold", fontSize=NAME_SIZE,
                               leading=NAME_SIZE + 2, alignment=TA_CENTER,
                               textColor="#000000", spaceAfter=2.6),
        "contact": ParagraphStyle("c", fontName=FONT, fontSize=CONTACT_SIZE,
                                  leading=CONTACT_SIZE + 2.2, alignment=TA_CENTER,
                                  textColor="#000000"),
        "block": ParagraphStyle("b", alignment=TA_LEFT, **body),
        "subject": ParagraphStyle("s", fontName=FONT + "-Bold", fontSize=BODY,
                                  leading=LEAD, alignment=TA_LEFT,
                                  textColor="#000000"),
        "para": ParagraphStyle("p", alignment=TA_JUSTIFY, spaceAfter=7.0, **body),
    }

    story = [P(NAME, st["name"])]
    parts = [f'<a href="{l}" color="black">{t}</a>' if l else t for t, l in CONTACT]
    story.append(P("  &#183;  ".join(parts), st["contact"]))
    story.append(Rule(FRAME_W, thickness=0.9, space_before=3.4))
    story.append(Spacer(1, 13))

    story.append(P(DATE, st["block"]))
    story.append(Spacer(1, 10))
    for line in RECIPIENT:
        story.append(P(line, st["block"]))
    story.append(Spacer(1, 12))
    story.append(P(SUBJECT, st["subject"]))
    story.append(Spacer(1, 12))
    story.append(P(SALUTATION, st["block"]))
    story.append(Spacer(1, 9))
    for para in PARAGRAPHS:
        story.append(P(para, st["para"]))
    story.append(Spacer(1, 5))
    story.append(P(CLOSING, st["block"]))
    story.append(Spacer(1, 16))
    story.append(P(f"<b>{SIGNOFF}</b>", st["block"]))
    story.append(Spacer(1, 8))
    story.append(P(ENCLOSURE, st["block"]))

    doc = BaseDocTemplate(
        path, pagesize=A4, leftMargin=MARGIN_X, rightMargin=MARGIN_X,
        topMargin=MARGIN_TOP, bottomMargin=MARGIN_BOTTOM,
        title="S K Lokesh - Cover Letter", author="S K Lokesh",
        subject=plain(SUBJECT))
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(
        MARGIN_X, MARGIN_BOTTOM, FRAME_W, A4[1] - MARGIN_TOP - MARGIN_BOTTOM,
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)])])
    doc.build(story)
    return path


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def build_docx(path=OUT_BASE + ".docx"):
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = R.DOCX_FONT
    style.font.size = Pt(BODY)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), R.DOCX_FONT)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.79)
    sec.top_margin = sec.bottom_margin = Inches(0.51)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p, 0, 2, line=NAME_SIZE + 2)
    r = p.add_run(NAME)
    r.font.name, r.font.size, r.bold = R.DOCX_FONT, Pt(NAME_SIZE), True

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p, 0, 3, line=CONTACT_SIZE + 2.6)
    for i, (text, link) in enumerate(CONTACT):
        if i:
            sep = p.add_run("  \u00b7  ")
            sep.font.name, sep.font.size = R.DOCX_FONT, Pt(CONTACT_SIZE)
        if link:
            _add_hyperlink(p, text, link)
        else:
            run = p.add_run(text)
            run.font.name, run.font.size = R.DOCX_FONT, Pt(CONTACT_SIZE)
    _bottom_border(p, size=8)

    def block(markup, before=0, after=0, bold=False):
        q = d.add_paragraph()
        _set_spacing(q, before, after, line=LEAD)
        _add_runs(q, f"<b>{markup}</b>" if bold else markup, BODY)
        return q

    block(DATE, before=13)
    for i, line in enumerate(RECIPIENT):
        block(line, before=10 if i == 0 else 0)
    block(SUBJECT, before=12, bold=True)
    block(SALUTATION, before=12)
    for i, para in enumerate(PARAGRAPHS):
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _set_spacing(q, 9 if i == 0 else 7, 0, line=LEAD)
        _add_runs(q, para, BODY)
    block(CLOSING, before=12)
    block(SIGNOFF, before=16, bold=True)
    block(ENCLOSURE, before=8)

    props = d.core_properties
    props.author, props.title = "S K Lokesh", "S K Lokesh - Cover Letter"
    props.subject = plain(SUBJECT)
    d.save(path)
    return path


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #
def build_txt(path=OUT_BASE + ".txt"):
    out = [plain(NAME), " | ".join(plain(t) for t, _ in CONTACT), "", DATE, ""]
    out += [plain(l) for l in RECIPIENT]
    out += ["", plain(SUBJECT), "", plain(SALUTATION), ""]
    for para in PARAGRAPHS:
        out += [plain(para), ""]
    out += [plain(CLOSING), "", plain(SIGNOFF), "", plain(ENCLOSURE)]
    text = "\n".join(out).rstrip() + "\n"
    for a, b in (("\u2013", "-"), ("\u2014", "-"), ("\u00b7", "|"), ("\u2019", "'")):
        text = text.replace(a, b)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
    return path


if __name__ == "__main__":
    for fn in (build_pdf, build_docx, build_txt):
        print("wrote", fn())
