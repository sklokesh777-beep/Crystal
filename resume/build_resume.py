#!/usr/bin/env python3
"""
Builds S K Lokesh's resume in three formats from one content source:

  S-K-Lokesh-Resume.pdf   - the file to upload to the JPMorganChase application
  S-K-Lokesh-Resume.docx  - editable Word master (same layout, Times New Roman)
  S-K-Lokesh-Resume.txt   - plain text for pasting into web forms / ATS checks

Design constraints (all deliberate):
  * single column, standard headings, no tables/text boxes/graphics carrying
    content -> parses cleanly in Oracle Recruiting Cloud, which JPMorganChase uses
  * one page, serif type (Liberation Serif = Times New Roman metrics)
  * dates right-aligned on the same physical line as the entry they belong to
  * no photo, no age, no address beyond city, no skill bars, no colour blocks

Usage:  python3 build_resume.py
"""

import html
import os
import re

from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, KeepTogether,
                               PageTemplate, Paragraph, Spacer, Table,
                               TableStyle)

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from resume_content import CONTACT, NAME, PROFILE, SECTIONS

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_BASE = os.path.join(HERE, "S-K-Lokesh-Resume")

# --------------------------------------------------------------------------- #
# typography
# --------------------------------------------------------------------------- #
FONT_DIR = "/usr/share/fonts/liberation-serif"
FONT = "Serif"

NAME_SIZE = 19.0
CONTACT_SIZE = 9.0
HEAD_SIZE = 10.2
BODY_SIZE = 9.5
BODY_LEAD = 11.35
SUB_SIZE = 8.9

INK = "#000000"
RULE = "#000000"

PAGE_W, PAGE_H = A4
MARGIN_X = 14.0 * mm
MARGIN_TOP = 11.0 * mm
MARGIN_BOTTOM = 9.0 * mm
FRAME_W = PAGE_W - 2 * MARGIN_X

TAG_RE = re.compile(r"<[^>]+>")


def register_fonts() -> None:
    faces = {
        FONT: "LiberationSerif-Regular.ttf",
        FONT + "-Bold": "LiberationSerif-Bold.ttf",
        FONT + "-Italic": "LiberationSerif-Italic.ttf",
        FONT + "-BoldItalic": "LiberationSerif-BoldItalic.ttf",
    }
    for name, filename in faces.items():
        pdfmetrics.registerFont(TTFont(name, os.path.join(FONT_DIR, filename)))
    pdfmetrics.registerFontFamily(
        FONT, normal=FONT, bold=FONT + "-Bold",
        italic=FONT + "-Italic", boldItalic=FONT + "-BoldItalic")


# --------------------------------------------------------------------------- #
# helpers shared by all three writers
# --------------------------------------------------------------------------- #
def smart(text: str) -> str:
    """Typographic apostrophe for the PDF/DOCX renderings (ASCII kept in .txt)."""
    return text.replace("'", "\u2019")


def plain(markup: str) -> str:
    """Strip inline markup and resolve entities -> clean text."""
    return html.unescape(TAG_RE.sub("", markup)).replace("  ", " ").strip()


def runs(markup: str):
    """Split '<b>bold</b> normal' into [(text, is_bold, is_italic), ...]."""
    markup = smart(markup)
    out, pos = [], 0
    pattern = re.compile(r"<(/?)(b|i)>")
    bold = italic = False
    for m in pattern.finditer(markup):
        chunk = markup[pos:m.start()]
        if chunk:
            out.append((html.unescape(chunk), bold, italic))
        closing, tag = m.group(1), m.group(2)
        if tag == "b":
            bold = not closing
        else:
            italic = not closing
        pos = m.end()
    tail = markup[pos:]
    if tail:
        out.append((html.unescape(tail), bold, italic))
    return out


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
class Rule(Flowable):
    """Hairline rule used under the name block and under each section heading."""

    def __init__(self, width, thickness=0.6, space_before=1.2, space_after=0.0):
        super().__init__()
        self.width = width
        self.thickness = thickness
        self.space_before = space_before
        self.space_after = space_after
        self.height = thickness + space_before + space_after

    def draw(self):
        self.canv.setLineWidth(self.thickness)
        self.canv.setStrokeColor(RULE)
        y = self.space_after + self.thickness / 2.0
        self.canv.line(0, y, self.width, y)


def P(markup, style, **kw):
    return Paragraph(smart(markup), style, **kw)


def pdf_styles():
    base = dict(fontName=FONT, textColor=INK, fontSize=BODY_SIZE, leading=BODY_LEAD)
    return {
        "name": ParagraphStyle("name", fontName=FONT + "-Bold", fontSize=NAME_SIZE,
                               leading=NAME_SIZE + 2, alignment=TA_CENTER,
                               textColor=INK, spaceAfter=2.6),
        "contact": ParagraphStyle("contact", fontName=FONT, fontSize=CONTACT_SIZE,
                                  leading=CONTACT_SIZE + 2.2, alignment=TA_CENTER,
                                  textColor=INK),
        "head": ParagraphStyle("head", fontName=FONT + "-Bold", fontSize=HEAD_SIZE,
                               leading=HEAD_SIZE + 1.5, textColor=INK,
                               spaceBefore=6.0, spaceAfter=1.0),
        "profile": ParagraphStyle("profile", alignment=TA_JUSTIFY, **base),
        "left": ParagraphStyle("left", alignment=TA_LEFT, **base),
        "right": ParagraphStyle("right", alignment=TA_RIGHT, **base),
        "sub": ParagraphStyle("sub", fontName=FONT + "-Italic", fontSize=SUB_SIZE,
                              leading=SUB_SIZE + 2.0, textColor=INK),
        "line": ParagraphStyle("line", alignment=TA_LEFT, **base),
        "bullet": ParagraphStyle("bullet", alignment=TA_LEFT, fontName=FONT,
                                 fontSize=BODY_SIZE, leading=BODY_LEAD,
                                 textColor=INK, leftIndent=10.5, bulletIndent=1.0,
                                 bulletFontName=FONT, bulletFontSize=BODY_SIZE,
                                 spaceBefore=1.4),
        "kv": ParagraphStyle("kv", alignment=TA_LEFT, fontName=FONT,
                             fontSize=BODY_SIZE, leading=BODY_LEAD, textColor=INK,
                             leftIndent=0, spaceBefore=1.6),
    }


def entry_row(left_markup, right_markup, st):
    """Left text + right-aligned date on one baseline (invisible 2-col layout)."""
    right_w = 0.0
    if right_markup:
        right_w = pdfmetrics.stringWidth(
            plain(right_markup), FONT + "-Bold", BODY_SIZE) + 6.0
    left_w = FRAME_W - right_w
    cells = [[P(left_markup, st["left"]),
              P(right_markup or "", st["right"])]]
    t = Table(cells, colWidths=[left_w, right_w])
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return t


def build_pdf(path=OUT_BASE + ".pdf"):
    register_fonts()
    st = pdf_styles()
    story = []

    # ---- header
    story.append(P(NAME, st["name"]))
    parts = []
    for text, link in CONTACT:
        parts.append(f'<a href="{link}" color="black">{text}</a>' if link else text)
    story.append(P("  &#183;  ".join(parts), st["contact"]))
    story.append(Rule(FRAME_W, thickness=0.9, space_before=3.4, space_after=0.0))

    # ---- profile (no heading label: the page opens on evidence, not a label)
    story.append(Spacer(1, 4.4))
    story.append(P(PROFILE, st["profile"]))

    # ---- sections
    for section in SECTIONS:
        head = [P(section["heading"].upper(), st["head"]),
                Rule(FRAME_W, thickness=0.6, space_before=0.6, space_after=0.0),
                Spacer(1, 2.6)]
        story.extend(head)

        for block in section["blocks"]:
            group = []
            if block.get("left") is not None:
                group.append(entry_row(block["left"], block.get("right"), st))
            if block.get("sub"):
                group.append(Spacer(1, 0.8))
                group.append(P(block["sub"], st["sub"]))
            if block["type"] != "kvlines":
                for line in block.get("lines", []):
                    group.append(Spacer(1, 1.4))
                    group.append(P(line, st["line"]))
            for bullet in block.get("bullets", []):
                group.append(P(bullet, st["bullet"], bulletText="\u2022"))
            if block["type"] == "kvlines":
                for label, value in block["lines"]:
                    group.append(P(f"<b>{label}:</b> {value}", st["kv"]))
            story.append(KeepTogether(group))
            if block is not section["blocks"][-1]:
                story.append(Spacer(1, 3.6))

    doc = BaseDocTemplate(
        path, pagesize=A4,
        leftMargin=MARGIN_X, rightMargin=MARGIN_X,
        topMargin=MARGIN_TOP, bottomMargin=MARGIN_BOTTOM,
        title="S K Lokesh - Resume",
        author="S K Lokesh",
        subject="JPMorganChase 2027 CIB Research & Analytics, Securities Services - Bengaluru",
        keywords=("fund accounting, portfolio accounting, NAV, financial reporting, reconciliations, "
                  "investor reporting, capital calls, distributions, alternative fund services, "
                  "data analytics, Python, Excel, Bloomberg, CFA Level I"),
    )
    frame = Frame(MARGIN_X, MARGIN_BOTTOM, FRAME_W,
                  PAGE_H - MARGIN_TOP - MARGIN_BOTTOM, id="body",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame])])
    doc.build(story)
    return path


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
DOCX_FONT = "Times New Roman"


def _set_spacing(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)


def _add_runs(p, markup, size=BODY_SIZE, italic_all=False):
    for text, bold, italic in runs(markup):
        r = p.add_run(text)
        r.font.name = DOCX_FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic or italic_all
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def _add_hyperlink(p, text, url, size=CONTACT_SIZE):
    part = p.part
    r_id = part.relate_to(
        url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:sz", str(int(size * 2))),
                     ("w:color", "000000")):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), DOCX_FONT)
            el.set(qn("w:hAnsi"), DOCX_FONT)
        else:
            el.set(qn("w:val"), val)
        rPr.append(el)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    p._p.append(link)


def _bottom_border(p, size=6):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))       # eighths of a point
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    pPr.append(borders)


def build_docx(path=OUT_BASE + ".docx"):
    d = docx.Document()

    style = d.styles["Normal"]
    style.font.name = DOCX_FONT
    style.font.size = Pt(BODY_SIZE)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    sec = d.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.55)
    sec.top_margin = Inches(0.43)
    sec.bottom_margin = Inches(0.35)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    def tabbed(markup_left, markup_right, size=BODY_SIZE, italic=False):
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
        _set_spacing(p, 0, 0, line=BODY_LEAD)
        _add_runs(p, markup_left, size, italic)
        if markup_right:
            p.add_run("\t")
            _add_runs(p, markup_right, size, italic)
        return p

    # ---- header
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p, 0, 2, line=NAME_SIZE + 2)
    r = p.add_run(NAME)
    r.font.name = DOCX_FONT
    r.font.size = Pt(NAME_SIZE)
    r.bold = True

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(p, 0, 3, line=CONTACT_SIZE + 2.6)
    for i, (text, link) in enumerate(CONTACT):
        if i:
            sep = p.add_run("  \u00b7  ")
            sep.font.name = DOCX_FONT
            sep.font.size = Pt(CONTACT_SIZE)
        if link:
            _add_hyperlink(p, text, link)
        else:
            r = p.add_run(text)
            r.font.name = DOCX_FONT
            r.font.size = Pt(CONTACT_SIZE)
    _bottom_border(p, size=8)

    # ---- profile
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_spacing(p, 4.4, 0, line=BODY_LEAD)
    _add_runs(p, PROFILE)

    # ---- sections
    for section in SECTIONS:
        h = d.add_paragraph()
        _set_spacing(h, 6.0, 2.6, line=HEAD_SIZE + 1.5)
        r = h.add_run(section["heading"].upper())
        r.font.name = DOCX_FONT
        r.font.size = Pt(HEAD_SIZE)
        r.bold = True
        _bottom_border(h, size=6)

        for idx, block in enumerate(section["blocks"]):
            if idx:
                spacer = d.add_paragraph()
                _set_spacing(spacer, 0, 0, line=3.6)
            if block.get("left") is not None:
                tabbed(block["left"], block.get("right"))
            if block.get("sub"):
                p = d.add_paragraph()
                _set_spacing(p, 0.8, 0, line=SUB_SIZE + 2.0)
                _add_runs(p, block["sub"], SUB_SIZE, italic_all=True)
            if block["type"] != "kvlines":
                for line in block.get("lines", []):
                    p = d.add_paragraph()
                    _set_spacing(p, 1.4, 0, line=BODY_LEAD)
                    _add_runs(p, line)
            for bullet in block.get("bullets", []):
                p = d.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Pt(11)
                pf.first_line_indent = Pt(-11)
                _set_spacing(p, 1.4, 0, line=BODY_LEAD)
                _add_runs(p, "\u2022   " + bullet)
            if block["type"] == "kvlines":
                for label, value in block["lines"]:
                    p = d.add_paragraph()
                    _set_spacing(p, 1.6, 0, line=BODY_LEAD)
                    _add_runs(p, f"<b>{label}:</b> {value}")

    props = d.core_properties
    props.author = "S K Lokesh"
    props.title = "S K Lokesh - Resume"
    props.subject = ("JPMorganChase 2027 CIB Research & Analytics, "
                     "Securities Services - Bengaluru")
    d.save(path)
    return path


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #
def build_txt(path=OUT_BASE + ".txt"):
    out = [plain(NAME), " | ".join(plain(t) for t, _ in CONTACT), ""]
    out += ["SUMMARY", plain(PROFILE), ""]
    for section in SECTIONS:
        out.append(plain(section["heading"]).upper())
        for block in section["blocks"]:
            if block.get("left") is not None:
                left, right = plain(block["left"]), plain(block.get("right") or "")
                out.append(f"{left}{('  |  ' + right) if right else ''}")
            if block.get("sub"):
                out.append(plain(block["sub"]))
            if block["type"] != "kvlines":
                for line in block.get("lines", []):
                    out.append(plain(line))
            for bullet in block.get("bullets", []):
                out.append("- " + plain(bullet))
            if block["type"] == "kvlines":
                for label, value in block["lines"]:
                    out.append(f"{plain(label)}: {plain(value)}")
            out.append("")
    text = "\n".join(out).rstrip() + "\n"
    text = (text.replace("\u2013", "-").replace("\u2014", "-")
                .replace("\u00b7", "|").replace("\u2019", "'")
                .replace("\u2022", "-"))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
    return path


if __name__ == "__main__":
    for fn in (build_pdf, build_docx, build_txt):
        print("wrote", fn())
